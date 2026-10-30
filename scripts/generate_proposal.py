from __future__ import annotations

from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TMP = ROOT / "tmp" / "documentacion"
OUT = ROOT / "entrega"
REF_PAGE = ROOT / "tmp" / "reference_render" / "page-1.png"
TMP.mkdir(parents=True, exist_ok=True)
OUT.mkdir(parents=True, exist_ok=True)

DOCX_PATH = OUT / "U2_Propuesta_Vigia_Suite.docx"
PURPLE = "6B50AD"
DARK = "252525"
MID = "5B5B5B"
LIGHT = "F3F0FA"
GRID = "D9D9D9"
WHITE = "FFFFFF"


def font_path(bold: bool = False, italic: bool = False) -> str:
    base = Path("C:/Windows/Fonts")
    if bold and italic:
        candidates = [base / "arialbi.ttf", base / "calibriz.ttf"]
    elif bold:
        candidates = [base / "arialbd.ttf", base / "calibrib.ttf"]
    elif italic:
        candidates = [base / "ariali.ttf", base / "calibrii.ttf"]
    else:
        candidates = [base / "arial.ttf", base / "calibri.ttf"]
    return str(next((p for p in candidates if p.exists()), candidates[0]))


def pil_font(size: int, bold: bool = False, italic: bool = False):
    return ImageFont.truetype(font_path(bold, italic), size)


def centered_text(draw, xy, text, font, fill=(0, 0, 0), spacing=8):
    box = draw.multiline_textbbox((0, 0), text, font=font, spacing=spacing, align="center")
    width = box[2] - box[0]
    draw.multiline_text((xy[0] - width / 2, xy[1]), text, font=font, fill=fill, spacing=spacing, align="center")


def build_cover() -> Path:
    width, height = 2550, 3300
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    bar_width = 315
    draw.rectangle((0, 0, bar_width, height), fill=(107, 80, 173))

    if REF_PAGE.exists():
        ref = Image.open(REF_PAGE).convert("RGB")
        # El recorte contiene únicamente el logotipo institucional del documento de referencia.
        logo = ref.crop((245, 70, 775, 310))
        logo.thumbnail((720, 380), Image.Resampling.LANCZOS)
        image.paste(logo, (bar_width + (width - bar_width - logo.width) // 2, 205))

    center_x = bar_width + (width - bar_width) / 2
    centered_text(draw, (center_x, 730), "UNIVERSIDAD POLITÉCNICA DE CHIAPAS", pil_font(62, bold=True))
    centered_text(draw, (center_x, 880), "Ingeniería en Tecnologías de la Información\ne Innovación Digital", pil_font(45), spacing=14)
    centered_text(draw, (center_x, 1370), "U2 - Propuesta de proyecto", pil_font(57, italic=True))
    centered_text(draw, (center_x, 1490), "Vigía Suite", pil_font(64, bold=True))
    centered_text(draw, (center_x, 1600), "Aplicación Android para diagnóstico defensivo de redes", pil_font(40, italic=True))

    centered_text(draw, (center_x, 1900), "Integrantes", pil_font(45, bold=True))
    centered_text(draw, (center_x, 2015), "Luis Alberto Nafate Hernández - 243728\nAlix Anahi Montesinos Grajales - 243777", pil_font(43), spacing=24)
    centered_text(draw, (center_x, 2290), "Grupo: 7B", pil_font(43, bold=True))
    centered_text(draw, (center_x, 2405), "Docente: Jose Alonso Macias Montoya", pil_font(42))
    centered_text(draw, (center_x, 2515), "Materia: Seguridad Informática", pil_font(42))
    centered_text(draw, (center_x, 2940), "SUCHIAPA, CHIAPAS, 8 DE OCTUBRE DE 2026", pil_font(42))

    path = TMP / "portada.png"
    image.save(path, quality=95)
    return path


def rounded_box(draw, box, fill, outline, title, subtitle="", title_fill="white", subtitle_fill="white"):
    draw.rounded_rectangle(box, radius=26, fill=fill, outline=outline, width=4)
    x1, y1, x2, y2 = box
    cx = (x1 + x2) / 2
    title_font = pil_font(31, bold=True)
    subtitle_font = pil_font(23)
    title_box = draw.textbbox((0, 0), title, font=title_font)
    draw.text((cx - (title_box[2] - title_box[0]) / 2, y1 + 28), title, font=title_font, fill=title_fill)
    if subtitle:
        lines = "\n".join(wrap(subtitle, 29))
        sb = draw.multiline_textbbox((0, 0), lines, font=subtitle_font, spacing=7, align="center")
        draw.multiline_text((cx - (sb[2] - sb[0]) / 2, y1 + 82), lines, font=subtitle_font, fill=subtitle_fill, spacing=7, align="center")


def arrow(draw, start, end, label="", color=(75, 75, 75)):
    draw.line((start, end), fill=color, width=6)
    x1, y1 = start
    x2, y2 = end
    import math
    angle = math.atan2(y2 - y1, x2 - x1)
    length = 18
    for offset in (2.55, -2.55):
        px = x2 + length * math.cos(angle + offset)
        py = y2 + length * math.sin(angle + offset)
        draw.line((x2, y2, px, py), fill=color, width=6)
    if label:
        f = pil_font(22, bold=True)
        bx = (x1 + x2) / 2
        by = (y1 + y2) / 2 - 34
        tb = draw.textbbox((0, 0), label, font=f)
        draw.rounded_rectangle((bx - (tb[2]-tb[0])/2 - 10, by - 4, bx + (tb[2]-tb[0])/2 + 10, by + 30), radius=8, fill="white")
        draw.text((bx - (tb[2]-tb[0])/2, by), label, font=f, fill=color)


def diagram_architecture() -> Path:
    im = Image.new("RGB", (1800, 1040), "white")
    d = ImageDraw.Draw(im)
    d.text((80, 55), "Arquitectura propuesta de Vigía Suite", font=pil_font(44, bold=True), fill=(0, 0, 0))
    d.text((80, 115), "Android es el producto obligatorio; la consulta de vulnerabilidades usa la API pública de NIST.", font=pil_font(27), fill=(70, 70, 70))

    rounded_box(d, (90, 300, 420, 540), "#6B50AD", "#584094", "ESP32 NodeMCU", "Escaneo Wi-Fi 2.4 GHz\nEstado del sensor\nFirmware Arduino C++")
    rounded_box(d, (600, 220, 980, 500), "#333333", "#111111", "Aplicación Android", "Kotlin y Jetpack Compose\nUSB OTG como canal principal\nAnálisis y presentación")
    rounded_box(d, (1190, 220, 1580, 500), "#6B50AD", "#584094", "NVD CVE API 2.0", "Servicio público de NIST\nBúsqueda por producto o CPE\nResultados CVE y CVSS")
    rounded_box(d, (600, 620, 980, 850), "#555555", "#333333", "Datos locales", "Room sobre SQLite\nHistorial de evaluaciones\nCaché de consultas")
    rounded_box(d, (1190, 640, 1580, 840), "#8A8A8A", "#666666", "Panel web opcional", "Solo si el Android termina antes\nNo forma parte del MVP")

    arrow(d, (420, 410), (600, 360), "USB serial")
    arrow(d, (980, 350), (1190, 350), "HTTPS GET")
    arrow(d, (790, 500), (790, 620), "Room")

    d.rounded_rectangle((90, 885, 1735, 980), radius=20, fill="#F3F0FA", outline="#CFC4E7", width=3)
    d.text((125, 915), "Límite defensivo: inventario y diagnóstico en redes propias o autorizadas. No se incluyen ataques, interferencia ni captura de credenciales.", font=pil_font(25, bold=True), fill=(74, 54, 122))
    path = TMP / "diagrama_arquitectura.png"
    im.save(path)
    return path


def diagram_sequence() -> Path:
    im = Image.new("RGB", (1800, 1150), "white")
    d = ImageDraw.Draw(im)
    d.text((80, 45), "Secuencia de una evaluación Wi-Fi", font=pil_font(44, bold=True), fill=(0, 0, 0))
    actors = [(190, "Usuario"), (600, "Android"), (1010, "ESP32"), (1440, "NVD API")]
    for x, name in actors:
        d.rounded_rectangle((x-120, 140, x+120, 220), radius=16, fill="#6B50AD", outline="#584094", width=3)
        tb = d.textbbox((0,0), name, font=pil_font(27, bold=True))
        d.text((x-(tb[2]-tb[0])/2, 162), name, font=pil_font(27, bold=True), fill="white")
        d.line((x, 220, x, 1070), fill="#B9B9B9", width=3)

    steps = [
        (290, 190, 600, "1  Conectar sensor"),
        (380, 600, 1010, "2  PING e INFO"),
        (470, 1010, 600, "3  PONG e identidad"),
        (560, 190, 600, "4  Iniciar Radar Wi-Fi"),
        (650, 600, 1010, "5  SCAN_WIFI"),
        (740, 1010, 600, "6  Resultados JSON Lines"),
        (830, 190, 600, "7  Consultar producto"),
        (900, 600, 1440, "8  GET /rest/json/cves/2.0"),
        (970, 1440, 600, "9  CVE y CVSS en JSON"),
        (1040, 600, 190, "10 Mostrar y guardar local"),
    ]
    for y, x1, x2, label in steps:
        arrow(d, (x1, y), (x2, y), label)
    path = TMP / "diagrama_secuencia.png"
    im.save(path)
    return path


def diagram_trust() -> Path:
    im = Image.new("RGB", (1800, 1000), "white")
    d = ImageDraw.Draw(im)
    d.text((80, 45), "Límites de confianza y controles", font=pil_font(44, bold=True), fill=(0, 0, 0))
    zones = [
        ((80, 180, 500, 850), "Zona física", "ESP32 y cable USB", "Firmware limitado a comandos permitidos\nIdentidad del sensor\nSin contraseñas Wi-Fi"),
        ((560, 180, 980, 850), "Zona Android", "Aplicación móvil", "Permiso explícito de puerto\nValidación de mensajes JSON\nConsultas elegidas por el usuario"),
        ((1040, 180, 1460, 850), "Zona de Internet", "NVD CVE API 2.0", "TLS\nLímite de frecuencia\nCaché y manejo de errores"),
        ((1520, 180, 1750, 850), "Datos", "Room local", "Cifrado del dispositivo\nSin contraseñas Wi-Fi\nBorrado controlado"),
    ]
    for idx, (box, title, subtitle, controls) in enumerate(zones):
        fill = "#F3F0FA" if idx % 2 == 0 else "#F5F5F5"
        d.rounded_rectangle(box, radius=24, fill=fill, outline="#6B50AD" if idx < 3 else "#555555", width=4)
        x1,y1,x2,y2=box
        d.text((x1+28,y1+30), title, font=pil_font(33,bold=True), fill=(0,0,0))
        d.text((x1+28,y1+86), subtitle, font=pil_font(25,italic=True), fill=(70,70,70))
        y=y1+175
        max_chars = max(13, int((x2 - x1) / 13))
        for line in controls.split("\n"):
            d.ellipse((x1+30,y+7,x1+44,y+21),fill="#6B50AD")
            wrapped=wrap(line,max_chars)
            d.multiline_text((x1+58,y),"\n".join(wrapped),font=pil_font(24),fill=(35,35,35),spacing=5)
            y += 62 + 28*(len(wrapped)-1)
    arrow(d,(500,520),(560,520),"UART")
    arrow(d,(980,520),(1040,520),"HTTPS")
    arrow(d,(1460,520),(1520,520),"SQL")
    path = TMP / "diagrama_confianza.png"
    im.save(path)
    return path


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def style_table(table, widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for r_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[c_idx])
            if r_idx == 0:
                set_cell_shading(cell, PURPLE)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                if r_idx == 0:
                    p.paragraph_format.keep_with_next = True
                for run in p.runs:
                    run.font.name = "Arial"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
                    run.font.size = Pt(9.5)
                    if r_idx == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, text in enumerate(headers):
        table.rows[0].cells[i].text = text
    for row_data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            cells[i].text = text
    style_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def set_run_font(run, name="Arial", size=12, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullets(doc, items):
    for index, item in enumerate(items):
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.space_before = Pt(3 if index == 0 else 0)
        p.paragraph_format.line_spacing = 1.1
        for run in p.runs:
            set_run_font(run)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.paragraph_format.space_after = Pt(5)
        for run in p.runs:
            set_run_font(run)


def add_figure(doc, image_path, caption, width=6.35):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(image_path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    r = cap.add_run(caption)
    set_run_font(r, size=9.5, italic=True, color=MID)


def set_page_break_before(paragraph):
    paragraph.paragraph_format.page_break_before = True


def add_heading(doc, text, level=1, page_break=False):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    if page_break:
        set_page_break_before(p)
    return p


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size in (("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 11.5)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(11 if style_name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def build_document():
    cover = build_cover()
    architecture = diagram_architecture()
    sequence = diagram_sequence()
    trust = diagram_trust()

    doc = Document()
    configure_styles(doc)
    first = doc.sections[0]
    first.page_width = Inches(8.5)
    first.page_height = Inches(11)
    first.top_margin = Inches(0)
    first.bottom_margin = Inches(0)
    first.left_margin = Inches(0)
    first.right_margin = Inches(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(1)
    p.add_run().add_picture(str(cover), width=Inches(8.5), height=Inches(10.88))

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.35)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Propuesta de proyecto Vigía Suite")
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    intro.paragraph_format.space_after = Pt(10)
    intro.paragraph_format.line_spacing = 1.15
    r = intro.add_run(
        "Vigía Suite será una solución portátil para revisar el entorno inalámbrico de una red propia o autorizada. "
        "Un ESP32 NodeMCU funcionará como sensor de proximidad y enviará observaciones a una aplicación Android. "
        "La aplicación conservará el historial localmente y consumirá la API REST NVD CVE 2.0 de NIST para consultar vulnerabilidades conocidas."
    )
    set_run_font(r)

    add_heading(doc, "1 Nombre del proyecto", 1)
    add_body(doc, "Vigía Suite Plataforma portátil para diagnóstico defensivo de redes")
    add_body(doc,
        "El nombre resume el propósito de la propuesta: reunir en una sola herramienta la conexión con el sensor, el inventario Wi-Fi, "
        "la consulta del estado del dispositivo y el registro de resultados. La palabra Vigía comunica observación preventiva sin presentar "
        "la aplicación como una herramienta de ataque.")

    add_heading(doc, "2 Planteamiento del problema a resolver", 1)
    add_heading(doc, "2.1 Contexto actual", 2)
    add_body(doc,
        "En una vivienda, aula o pequeño negocio es común encontrar teléfonos, televisores, computadoras, cámaras y dispositivos IoT conectados "
        "a la misma infraestructura inalámbrica. El propietario suele conocer el nombre de su red, pero no conserva un inventario ni puede "
        "comparar con facilidad qué redes aparecen alrededor, qué canales utilizan o cuáles anuncian mecanismos de autenticación débiles.")
    add_body(doc,
        "Las herramientas profesionales permiten obtener mucha información, aunque con frecuencia exigen una computadora, controladores especiales "
        "y conocimientos técnicos. En el otro extremo, varias aplicaciones móviles muestran datos aislados y dependen únicamente de las restricciones "
        "del teléfono. Esto deja un espacio para una solución educativa de bajo costo que separe la medición del dispositivo de visualización y explique "
        "los resultados en términos sencillos.")

    add_heading(doc, "2.2 Problema central", 2)
    add_body(doc,
        "El problema que atenderá Vigía Suite es la falta de una forma accesible, portátil y repetible de observar el entorno Wi-Fi y conservar evidencia "
        "de sus cambios. Sin un historial resulta difícil responder preguntas básicas: si apareció una red abierta cerca del lugar, si cambió el canal del "
        "punto de acceso, si el sensor dejó de funcionar o si un resultado pertenece a una evaluación anterior.")

    add_heading(doc, "2.3 Relevancia e impacto esperado", 2)
    add_body(doc,
        "Resolver esta necesidad ayudará a que un usuario sin experiencia en análisis de radio pueda realizar una revisión inicial y reconocer situaciones "
        "que merecen atención. La solución no sustituirá una auditoría profesional ni prometerá identificar por sí sola a un atacante. Su valor estará en "
        "hacer visibles señales básicas, conservarlas y orientar la siguiente acción.")
    add_bullets(doc, [
        "Reducir el tiempo necesario para obtener un inventario inicial del entorno Wi-Fi.",
        "Detectar redes abiertas, señales desconocidas o cambios entre evaluaciones.",
        "Separar la observación del ESP32 de la presentación de resultados en la aplicación.",
        "Mantener la información bajo control del usuario con un enfoque local y permisos explícitos.",
    ])

    add_heading(doc, "2.4 Caso de uso personal", 2)
    add_body(doc,
        "El caso principal será revisar una red doméstica o el laboratorio del alumno. El usuario conectará el ESP32 a su teléfono Android mediante USB OTG, "
        "validará la identidad del sensor y ejecutará Radar Wi-Fi. La aplicación ordenará los puntos de acceso por intensidad, mostrará canal, BSSID y tipo "
        "de autenticación, y marcará como prioritaria una red abierta. Si el usuario conoce el fabricante o modelo de un equipo, podrá consultar vulnerabilidades "
        "publicadas en NVD. El historial se guardará dentro del teléfono para compararlo con revisiones posteriores.")

    add_heading(doc, "3 Objetivos", 1)
    add_heading(doc, "3.1 Objetivo general", 2)
    add_body(doc,
        "Diseñar e implementar una aplicación Android integrada con un ESP32 NodeMCU y la API REST NVD CVE 2.0, capaz de recopilar, presentar y conservar "
        "observaciones defensivas de redes Wi-Fi autorizadas y consultar vulnerabilidades conocidas de productos identificados por el usuario.")

    add_heading(doc, "3.2 Objetivos específicos", 2)
    add_numbered(doc, [
        "Definir una arquitectura con límites claros entre el ESP32, la aplicación Android, la API pública y el almacenamiento local, incluyendo las amenazas principales y sus controles.",
        "Desarrollar el firmware del ESP32 para responder a comandos permitidos, reportar su estado y producir hasta 32 observaciones Wi-Fi por evaluación.",
        "Implementar una interfaz Android que gestione el permiso USB, valide la comunicación y presente resultados comprensibles sin requerir acceso root.",
        "Consumir la API NVD CVE 2.0 mediante HTTPS para buscar vulnerabilidades por producto, CPE o identificador CVE y manejar límites o fallos del servicio.",
        "Realizar pruebas funcionales, de integración y de seguridad sobre la conexión serial, el almacenamiento local y el consumo de la API.",
    ])

    add_heading(doc, "3.3 Alcance del producto mínimo viable", 2)
    add_table(doc,
        ["Incluido en el prototipo", "Fuera del alcance inicial"],
        [
            ("Conexión USB serial con CP210x o CH340", "Interferencia o desautenticación de redes"),
            ("Comandos PING INFO STATUS y SCAN_WIFI", "Captura de contraseñas o contenido privado"),
            ("Inventario de SSID BSSID canal señal y autenticación", "Compatibilidad con Wi-Fi de 5 GHz"),
            ("Historial local con Room y consultas a NVD", "Sustitución de una auditoría profesional"),
            ("Alertas por red abierta o cambio observado", "Explotación automática de vulnerabilidades"),
            ("Aplicación Android como entrega principal", "Panel web obligatorio para el MVP"),
        ], [3.2, 3.2])

    add_heading(doc, "4 Stack tecnológico", 1)
    add_body(doc,
        "El stack reduce el alcance a una aplicación Android. El ESP32 realiza la observación, Android presenta y conserva los resultados, y NVD aporta información "
        "pública de vulnerabilidades. El panel web quedará como una mejora opcional si el producto principal se termina antes de la fecha de entrega.")
    add_table(doc,
        ["Capa", "Tecnología", "Función y razón de elección"],
        [
            ("Aplicación móvil", "Kotlin y Jetpack Compose", "Interfaz Android nativa, gestión de permisos USB y conexión OTG sin requerir root."),
            ("Firmware", "Arduino C++ sobre ESP32", "Facilita la carga desde Arduino IDE y aprovecha Wi-Fi, BLE y UART del NodeMCU."),
            ("Consumo de API", "Retrofit OkHttp y Kotlin Coroutines", "Realiza solicitudes HTTPS, interpreta JSON y evita bloquear la interfaz."),
            ("API externa", "NIST NVD CVE API 2.0", "Servicio público para buscar CVE, CPE y métricas CVSS sin construir un servidor propio."),
            ("Datos locales", "Room sobre SQLite", "Mantiene evaluaciones, observaciones y consultas recientes dentro del teléfono."),
            ("Seguridad", "Android Keystore TLS y validación", "Protege configuración local, cifra el transporte y rechaza respuestas inesperadas."),
            ("Pruebas", "JUnit MockWebServer Android Lint", "Prueba el parser, simula respuestas de NVD y detecta errores comunes en Android."),
            ("Extensión opcional", "HTML CSS JavaScript Web Serial", "Solo se presentará si Android está terminado y probado antes de la entrega."),
        ], [1.25, 1.85, 3.3])

    add_heading(doc, "4.1 Decisión sobre el ESP32 y la antena", 2)
    add_body(doc,
        "Se utilizará el NodeMCU ESP-32S que ya fue adquirido. La placa integra Wi-Fi de 2.4 GHz y Bluetooth, y normalmente conecta el chip al USB mediante "
        "un convertidor CP210x o CH340. Esa decisión elimina una compra adicional y permite iniciar las pruebas de inmediato. Su antena impresa es suficiente "
        "para una vivienda, aula o laboratorio pequeño, siempre que la placa se mantenga alejada de metal y fuentes de ruido.")
    add_table(doc,
        ["Factor", "NodeMCU ESP-32S actual", "Opción con antena externa"],
        [
            ("Costo", "Sin costo adicional porque la placa ya está disponible", "Requiere otro módulo, conector y antena"),
            ("Cobertura", "Adecuada para pruebas cercanas y demostraciones", "Puede mejorar la estabilidad en áreas grandes"),
            ("Complejidad", "Carga directa desde Arduino IDE", "Exige seleccionar módulo compatible y montar la antena correctamente"),
            ("Decisión", "Elegido para el MVP", "Reservado para una segunda versión si las pruebas muestran pérdida de señal"),
        ], [1.25, 2.55, 2.6])
    add_body(doc,
        "La antena externa no se adoptará por anticipado. Primero se medirá el desempeño real del equipo actual. Si la cobertura resulta insuficiente, una "
        "segunda versión podrá usar un módulo ESP32-WROOM-32U con conector IPEX. Conectar una antena externa a una placa que no esté diseñada para ella puede "
        "empeorar la señal o dañar la etapa de radio.")

    add_heading(doc, "4.2 Riesgos técnicos del hardware y firmware", 2)
    add_bullets(doc, [
        "El ESP-32S trabaja únicamente en Wi-Fi de 2.4 GHz, por lo que una red exclusiva de 5 GHz no aparecerá en el inventario.",
        "Wi-Fi y Bluetooth comparten el radio de 2.4 GHz. Ejecutarlos al mismo tiempo puede reducir el rendimiento; el MVP usará USB serial como canal principal.",
        "El convertidor USB de la placa no es USB nativo. La aplicación debe reconocer CP210x o CH340 y manejar el permiso del sistema operativo.",
        "Una actualización defectuosa puede dejar el sensor sin responder. El firmware mostrará versión, identidad y estado antes de iniciar cada evaluación.",
        "Secure Boot y cifrado de flash existen en la familia ESP32, pero activar eFuses es una decisión delicada. En el proyecto académico primero se verificará la integridad del archivo con SHA-256 y se documentará el proceso de actualización.",
    ])

    add_heading(doc, "5 Arquitectura e integración", 1)
    add_figure(doc, architecture, "Figura 1. Componentes principales y canales de comunicación.")
    add_body(doc,
        "El ESP32 no almacenará cuentas ni contraseñas. Su responsabilidad será medir y devolver eventos en formato JSON Lines. Android validará esos eventos, "
        "los presentará y conservará el historial mediante Room. Cuando el usuario seleccione un fabricante, producto o CPE, la aplicación consultará NVD mediante "
        "HTTPS. La consulta será explícita porque un SSID o un BSSID, por sí solos, no demuestran qué producto existe ni que sea vulnerable.")

    add_heading(doc, "5.1 Protocolo entre la aplicación y el sensor", 2)
    add_table(doc,
        ["Comando", "Respuesta esperada", "Propósito"],
        [
            ("PING", "pong", "Comprobar que el enlace funciona antes de solicitar datos."),
            ("INFO", "identity", "Obtener placa, chip, identificador y versión del protocolo."),
            ("STATUS", "status", "Consultar memoria libre y tiempo de actividad."),
            ("SCAN_WIFI", "wifi_scan_started wifi_ap wifi_scan_finished", "Realizar el inventario limitado a 32 redes y cerrar la operación de forma explícita."),
        ], [1.1, 2.4, 2.9])
    add_body(doc,
        "El firmware rechazará cualquier orden que no aparezca en esta lista. La longitud del comando y el número de resultados también estarán limitados para "
        "reducir errores por entradas inesperadas. Los mensajes de arranque que no sean JSON se mostrarán en la consola de diagnóstico, pero no se guardarán como resultados.")

    add_heading(doc, "5.2 Flujo de una evaluación", 2)
    add_figure(doc, sequence, "Figura 2. Interacción desde la conexión del sensor hasta el informe.")

    add_heading(doc, "6 Integración con la API REST NVD", 1)
    add_body(doc,
        "Se consumirá la API NVD CVE 2.0 del National Institute of Standards and Technology. Es un servicio público y gratuito, ampliamente reconocido para consultar vulnerabilidades conocidas, y devuelve "
        "JSON. Para el prototipo se realizarán pocas consultas y se conservará temporalmente la respuesta para no repetir solicitudes. Una clave de API de NVD será "
        "opcional y, si se utiliza, no se incluirá directamente en el código fuente.")
    add_table(doc,
        ["Método y ruta de NVD", "Uso en Android", "Tratamiento"],
        [
            ("GET /rest/json/cves/2.0?keywordSearch={texto}", "Buscar CVE por fabricante o producto", "Codificar parámetros, limitar resultados y paginar"),
            ("GET /rest/json/cves/2.0?cveId={id}", "Consultar un identificador CVE concreto", "Validar el formato CVE antes de enviar"),
            ("GET /rest/json/cpes/2.0?keywordSearch={producto}", "Encontrar nombres CPE candidatos", "El usuario confirma el producto correcto"),
            ("GET /rest/json/cves/2.0?cpeName={cpe}", "Obtener vulnerabilidades asociadas a un CPE", "Mostrar versión y fecha de la información"),
        ], [2.25, 2.25, 2.0])

    add_heading(doc, "6.1 Datos utilizados", 2)
    add_body(doc,
        "Room almacenará Evaluación, Observación de red, Consulta de producto y Vulnerabilidad consultada. Cada observación guardará SSID, BSSID, canal, RSSI, "
        "autenticación y fecha. De NVD se conservarán únicamente campos útiles como identificador CVE, descripción, fecha, referencias y puntuación CVSS disponible. "
        "Las contraseñas Wi-Fi no forman parte del modelo.")

    add_heading(doc, "6.2 Ejemplo de integración", 2)
    add_body(doc,
        "Después del escaneo, el usuario podrá escribir o seleccionar el producto que desea investigar. Android buscará primero un CPE candidato y pedirá confirmación "
        "antes de consultar vulnerabilidades. La aplicación mostrará que una coincidencia de texto no confirma que el equipo observado tenga esa versión. Si NVD no "
        "responde o limita la solicitud, se conservará la evaluación local y se permitirá reintentar después.")

    add_heading(doc, "7 Seguridad por diseño", 1)
    add_figure(doc, trust, "Figura 3. Límites de confianza y controles previstos.")
    add_body(doc,
        "La lógica adversaria parte de asumir que el cable puede desconectarse, el sensor puede enviar datos dañados, la conexión a Internet puede fallar y una búsqueda "
        "por nombre puede asociar el producto equivocado. Por ello, Android validará cada frontera y distinguirá una observación local de una vulnerabilidad confirmada.")
    add_table(doc,
        ["Activo", "Escenario adversario", "Control previsto"],
        [
            ("Puerto serial", "Otro programa ocupa el puerto o un dispositivo distinto responde", "Permiso explícito, filtro por VID y validación de identidad y protocolo"),
            ("Firmware", "Comando largo o desconocido provoca comportamiento inesperado", "Lista permitida, búfer limitado y respuesta de error"),
            ("Resultados", "El sensor envía JSON incompleto o un número excesivo de redes", "Validación estricta, máximo de 32 observaciones y rechazo de campos fuera de rango"),
            ("Consulta NVD", "El nombre coincide con otro producto o versión", "Confirmación del usuario, mostrar CPE y no afirmar vulnerabilidad automática"),
            ("API externa", "Servicio lento, sin conexión o límite de solicitudes", "Tiempo de espera, caché, reintento controlado y manejo de HTTP 429"),
            ("Datos locales", "Otra aplicación o un respaldo expone el historial", "Room privado, cifrado del dispositivo y borrado desde la interfaz"),
            ("Actualización", "Carga de firmware modificado", "Distribución controlada y verificación SHA-256 antes de cargar"),
        ], [1.25, 2.7, 2.55])

    add_heading(doc, "7.1 Privacidad y uso responsable", 2)
    add_body(doc,
        "La interfaz explicará qué datos se capturan antes de comenzar. Un SSID puede revelar nombres personales o de negocios, por lo que el usuario podrá borrar "
        "una evaluación y definir cuánto tiempo conservarla. Las consultas a NVD contendrán términos de producto o CPE, no SSID ni BSSID. La aplicación no intentará "
        "conectarse a las redes encontradas, no guardará contraseñas y no incluirá funciones de interferencia, suplantación o desautenticación.")

    add_heading(doc, "8 Plan de implementación y pruebas", 1)
    add_heading(doc, "8.1 Estado actual", 2)
    add_table(doc,
        ["Componente", "Estado", "Evidencia funcional"],
        [
            ("Firmware Arduino", "Prototipo disponible", "Responde a PING INFO STATUS y SCAN_WIFI"),
            ("Aplicación Android", "Código base disponible", "Detecta CP210x y CH340 y abre 115200 8N1"),
            ("Integración NVD", "Pendiente", "Se consumirá CVE API 2.0 mediante Retrofit"),
            ("Room local", "Pendiente", "Guardará evaluaciones y caché de consultas"),
            ("Panel web", "Extensión opcional", "Solo se presentará si Android termina antes"),
        ], [1.8, 1.55, 3.15])

    add_heading(doc, "8.2 Fases de trabajo", 2)
    add_table(doc,
        ["Fase", "Producto verificable", "Criterio de salida"],
        [
            ("1 Diseño", "Arquitectura contratos y amenazas", "Revisión del alcance Android y rutas NVD"),
            ("2 Sensor", "Firmware estable", "Cien ciclos de conexión y escaneo sin bloqueo"),
            ("3 Android", "Interfaz y conexión USB", "Escaneo completo desde el teléfono"),
            ("4 API y datos", "NVD y Room", "Consulta CVE e historial local probados"),
            ("5 Seguridad", "Informe de pruebas", "Sin fallas críticas de autorización o validación"),
            ("6 Demostración", "Escenario completo", "Escaneo guardado y consulta NVD desde Android"),
        ], [0.9, 2.55, 3.05])

    add_heading(doc, "8.3 Pruebas propuestas", 2)
    add_bullets(doc, [
        "Pruebas funcionales de conexión, desconexión, reconexión y manejo de puerto ocupado.",
        "Pruebas de integración para comprobar que Android interpreta correctamente los eventos enviados por el firmware.",
        "Pruebas de límites con comandos largos, JSON dañado, campos fuera de rango y más de 32 observaciones.",
        "Pruebas con respuestas NVD válidas, vacías, incompletas, HTTP 429, pérdida de Internet y tiempo de espera.",
        "Pruebas de privacidad para confirmar que no se envían SSID, BSSID, contraseñas ni datos seriales a NVD.",
        "Prueba de aceptación en una red propia con una evaluación inicial, un cambio controlado y una segunda evaluación que genere una alerta.",
    ])

    add_heading(doc, "9 Factibilidad y criterios de evaluación", 1)
    add_body(doc,
        "La propuesta es factible porque el hardware principal ya está disponible y los componentes de software seleccionados son gratuitos. El MVP limita el número "
        "de comandos, resultados y consultas externas, lo que permite demostrar la integración completa sin desarrollar un servidor propio. Android es la única entrega "
        "obligatoria. El panel web se conservará como extensión opcional y no condicionará la evaluación del MVP.")
    add_table(doc,
        ["Criterio", "Respuesta de la propuesta"],
        [
            ("Innovación y relevancia 30 por ciento", "Sensor económico separado del cliente, operación local y explicación de riesgos para usuarios no especialistas."),
            ("Factibilidad técnica 30 por ciento", "ESP32 disponible, firmware acotado, Android como único cliente obligatorio y API pública de NIST."),
            ("Claridad 20 por ciento", "El problema, los objetivos, el alcance y cada tecnología responden al mismo caso de uso."),
            ("Integración completa 20 por ciento", "ESP32 y Android usan JSON Lines; Android consume NVD mediante JSON sobre HTTPS."),
        ], [2.25, 4.25])

    add_heading(doc, "10 Resultado esperado", 1)
    add_body(doc,
        "Al finalizar, una demostración completa deberá permitir conectar el NodeMCU ESP-32S, verificar su identidad, ejecutar una evaluación Wi-Fi autorizada, "
        "visualizar las redes por intensidad y tipo de autenticación, guardar la evaluación localmente y consultar en NVD un producto identificado por el usuario. "
        "El resultado mostrará CVE y CVSS cuando existan, permitirá eliminar el historial y cubrirá Android, hardware, API y controles de seguridad.")

    add_heading(doc, "11 Conclusión", 1)
    add_body(doc,
        "Vigía Suite propone una solución centrada en una necesidad real: entender el entorno inalámbrico sin depender de equipos costosos ni de una interfaz diseñada "
        "para especialistas. El ESP32 ya adquirido permite comenzar con un costo marginal bajo, mientras que Android, Room y NVD forman un alcance realizable. "
        "La propuesta reconoce las limitaciones del hardware, evita funciones ofensivas y deja el panel web únicamente como mejora posterior si el MVP se concluye antes.")

    add_heading(doc, "12 Bibliografía", 1)
    references = [
        "Espressif Systems. ESP32 Series Datasheet. https://www.espressif.com/sites/default/files/documentation/esp32_datasheet_en.pdf",
        "National Vulnerability Database. Vulnerability API. https://nvd.nist.gov/developers/vulnerabilities",
        "National Vulnerability Database. Product API. https://nvd.nist.gov/developers/products",
        "Android Developers. Save data in a local database using Room. https://developer.android.com/training/data-storage/room",
        "OWASP Foundation. OWASP API Security Top 10 2023. https://api-security.owasp.org/editions/2023/en/0x10-api-security-risks/",
        "National Institute of Standards and Technology. SP 800-115 Technical Guide to Information Security Testing and Assessment. https://csrc.nist.gov/pubs/sp/800/115/final",
    ]
    add_bullets(doc, references)

    props = doc.core_properties
    props.title = "Propuesta de proyecto Vigía Suite"
    props.subject = "Aplicación Android con ESP32 y NVD CVE API para diagnóstico defensivo de redes"
    props.author = "Luis Alberto Nafate Hernández y Alix Anahi Montesinos Grajales"
    props.keywords = "ciberseguridad, ESP32, Android, NVD, CVE, API REST, Room"
    props.comments = "Propuesta académica para la materia Seguridad Informática"

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build_document())
